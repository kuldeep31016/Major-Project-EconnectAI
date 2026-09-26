"""EcoConnectAI final presentation in the college-mandated format (major project final - ppt format.pptx).

Order (college format): Title · Abstract · Literature Survey ×7 · Problem & Objectives · Existing System ×2 ·
Proposed System ×4 · Results & Discussion ×4 · Conclusion · Co-guide meeting proof · SDG Goal ·
Project Stakeholders & Existing Application Gap (college template slide, filled) · References · Thank You.
All project numbers come from the repository's run/evaluation files.
"""
import copy
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from lxml import etree

SRC, OUT, IMG = "fmt.pptx", "EcoConnectAI_Final_Presentation_College_Format.pptx", "img/"
NAVY, ORANGE, PEACH = "0D3889", "D96A2B", "FBEFE6"
GREEN, MINT, LIGHT = "0E9E74", "E7F6F1", "EEF2F8"
TEXT, MUTED, WHITE, LINE, GOLD = "1A1A1A", "5A6370", "FFFFFF", "C9D3E6", "FFD54A"
FONT = "Calibri"
rgb = RGBColor.from_string

prs = Presentation(SRC)
# keep only the college "Project Stakeholders" template slide (3rd); drop the instruction/SDG/blank slides
sld = prs.slides._sldIdLst
ids = list(sld)
STAKE_ID = ids[2]
for s in ids:
    if s is not STAKE_ID:
        prs.part.drop_rel(s.rId); sld.remove(s)
stake_slide = prs.slides[0]
# rename the kept template slide's parts so newly added slides (slide1..N) cannot collide with it
from pptx.opc.packuri import PackURI
stake_slide.part.partname = PackURI("/ppt/slides/slide999.xml")
if stake_slide.has_notes_slide:
    stake_slide.notes_slide.part.partname = PackURI("/ppt/notesSlides/notesSlide999.xml")
LAYOUT = [l for l in prs.slide_layouts if l.name == "Title Only"][0]
order = []   # slide objects in final order


def new_slide(title, size=30):
    s = prs.slides.add_slide(LAYOUT)
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx != 0:
            ph._element.getparent().remove(ph._element)
    t = s.shapes.title
    t.text_frame.text = title
    p = t.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    for r in p.runs:
        r.font.size = Pt(size); r.font.name = FONT; r.font.color.rgb = rgb(TEXT); r.font.bold = False
    t.left, t.top, t.width, t.height = Inches(1.15), Inches(0.22), Inches(8.6), Inches(0.9)
    t.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    order.append(s)
    return s


def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt


def box(s, x, y, w, h, fill=WHITE, line=None, radius=0.12, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=1.0, dash=False):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw)
        if dash:
            from pptx.enum.dml import MSO_LINE
            sh.line.dash_style = MSO_LINE.DASH
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, paras, size=14, color=TEXT, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         italic=False, space=4):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04); tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space)
        for rt, o in (para if isinstance(para, list) else [(para, {})]):
            r = p.add_run(); r.text = rt; f = r.font
            f.name = FONT; f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic); f.color.rgb = rgb(o.get("color", color))
    return tb


def badge(s, x, y, d, label, fill, size=14, color=WHITE):
    c = box(s, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label; r.font.size = Pt(size); r.font.bold = True; r.font.name = FONT
    r.font.color.rgb = rgb(color)
    return c


def lbox(s, x, y, w, h, label, fill=WHITE, line=NAVY, color=NAVY, size=12.5, bold=True, sub=None, subsize=10,
         align=PP_ALIGN.CENTER):
    b = box(s, x, y, w, h, fill=fill, line=line)
    tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.08); tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = label; r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT
    r.font.color.rgb = rgb(color)
    if sub:
        p2 = tf.add_paragraph(); p2.alignment = align
        r2 = p2.add_run(); r2.text = sub; r2.font.size = Pt(subsize); r2.font.name = FONT
        r2.font.color.rgb = rgb(WHITE if fill in (NAVY, GREEN, ORANGE) else MUTED)
    return b


def arrow(s, x, y, w=0.28, h=0.22, direction="right", color=NAVY):
    shp = {"right": MSO_SHAPE.RIGHT_ARROW, "down": MSO_SHAPE.DOWN_ARROW, "left": MSO_SHAPE.LEFT_ARROW}[direction]
    return box(s, x, y, w, h, fill=color, shape=shp)


def picture(s, path, x, y, w=None, h=None):
    kw = {}
    if w: kw["width"] = Inches(w)
    if h: kw["height"] = Inches(h)
    pic = s.shapes.add_picture(IMG + path, Inches(x), Inches(y), **kw)
    pic.line.color.rgb = rgb(LINE); pic.line.width = Pt(1)
    return pic


def banner(s, x, y, w, h, runs, size=13, fill=NAVY):
    b = box(s, x, y, w, h, fill=fill, radius=0.18)
    tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.15)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    for rt, o in runs:
        r = p.add_run(); r.text = rt; r.font.size = Pt(o.get("size", size)); r.font.bold = o.get("bold", False)
        r.font.name = FONT; r.font.italic = o.get("italic", False); r.font.color.rgb = rgb(o.get("color", WHITE))
    return b


def bullets(s, x, y, w, items, dot=NAVY, size=12.5, gap=0.5, h=0.5):
    for it in items:
        box(s, x, y + 0.1, 0.09, 0.09, fill=dot, shape=MSO_SHAPE.OVAL)
        text(s, x + 0.18, y, w - 0.18, h, it if isinstance(it, list) else it, size=size)
        y += gap
    return y


def table(s, x, y, w, col_w, rows, header, size=10.5, head_size=11.5, row_h=None):
    tbl = s.shapes.add_table(len(rows) + 1, len(header), Inches(x), Inches(y), Inches(w), Inches(0.4)).table
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    def cell(c, txt, fill, color, bold, sz):
        c.fill.solid(); c.fill.fore_color.rgb = rgb(fill)
        c.margin_left = c.margin_right = Inches(0.07); c.margin_top = c.margin_bottom = Inches(0.05)
        tf = c.text_frame; tf.word_wrap = True
        paras = txt if isinstance(txt, list) else [txt]
        for i, ptxt in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            runs = ptxt if isinstance(ptxt, list) else [(ptxt, {})]
            for rt, o in runs:
                r = p.add_run(); r.text = rt; r.font.size = Pt(o.get("size", sz)); r.font.name = FONT
                r.font.bold = o.get("bold", bold); r.font.italic = o.get("italic", False)
                r.font.color.rgb = rgb(o.get("color", color))
    for j, hd in enumerate(header):
        cell(tbl.cell(0, j), hd, NAVY, WHITE, True, head_size)
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            cell(tbl.cell(i, j), v, LIGHT if i % 2 else WHITE, TEXT, False, size)
        if row_h:
            tbl.rows[i].height = Inches(row_h)
    tbl.rows[0].height = Inches(0.4)
    # plain grid look (no banding style)
    tblPr = tbl._tbl.tblPr
    tblPr.set("bandRow", "0"); tblPr.set("firstRow", "1")
    return tbl


# =============================================================================== 1 · TITLE
s = new_slide("EcoConnectAI", size=44)
text(s, 1.35, 1.15, 8.1, 0.9, [[("A Satellite-Driven Framework for Coastal Ecosystem Connectivity", {})],
                               [("and Conservation Decision Support", {})]],
     size=19, color=NAVY, italic=True, align=PP_ALIGN.CENTER, space=0)
text(s, 1.35, 2.35, 8.1, 1.2, [[("Presented by", {"size": 12, "color": MUTED})], [("Kuldeep Raj  (1DS23CS112)", {})],
                               [("Ruhinaaz  (1DS23CS187)", {})]], size=15, align=PP_ALIGN.CENTER, space=2)
text(s, 1.35, 3.75, 8.1, 0.75, [[("Guided by", {"size": 12, "color": MUTED})],
                                [("Dr. Anusha Preetham  (Associate Professor)", {})]], size=15, align=PP_ALIGN.CENTER, space=2)
text(s, 1.35, 4.75, 8.1, 0.9, ["Department of Computer Science and Engineering", "Dayananda Sagar College of Engineering",
                               "Final Year Project Evaluation · 2026"], size=12.5, color=MUTED, align=PP_ALIGN.CENTER, space=1)
notes(s, "Good morning. Our project is EcoConnectAI — a tool that uses satellite images and AI to find mangrove habitat "
         "and then works out which habitat patches matter most for keeping the coastal ecosystem connected.")

# =============================================================================== 2 · ABSTRACT
s = new_slide("Abstract")
cards = [("Problem overview", ORANGE, PEACH,
          "Mangrove maps show where habitat is, but not which patches hold the coastal network together. "
          "Losing one small linking patch can cut the network apart."),
         ("Proposed solution", NAVY, LIGHT,
          "Detect mangroves from Sentinel-1 satellite images with a deep-learning model, split them into patches, "
          "connect nearby patches into a network, and remove one patch at a time to see how much connectivity is lost. "
          "A web dashboard supports what-if and restoration questions."),
         ("Key results", GREEN, MINT,
          "Development model: test IoU 0.842, F1 0.914 against Global Mangrove Watch reference labels. "
          "Example: patch P17 is only 3.1 ha, yet removing it cuts connectivity by 27% — size alone does not show importance.")]
yy = 1.3
for h, c, f, body in cards:
    box(s, 1.3, yy, 8.3, 1.38, fill=f)
    text(s, 1.5, yy + 0.08, 7.9, 1.28, [[(h, {"bold": True, "color": c, "size": 15})], [(body, {"size": 13})]], space=3)
    yy += 1.5
banner(s, 1.3, 5.85, 8.3, 0.45, [("In short:  ", {"bold": True, "color": GOLD}),
                                  ("from “where are the mangroves?” to “which mangrove patch matters most?”", {})], size=12.5)
notes(s, "Problem: maps show where mangroves are, not which patches matter. Our solution: detect mangroves from radar "
         "images, turn them into a network of patches, and test what is lost if each patch disappears. Key result: our "
         "development model reaches IoU 0.842 against reference labels, and a 3.1-hectare patch turned out to be one of "
         "the most important in Kerala.")

# =============================================================================== 3–9 · LITERATURE SURVEY (7)
HDR = ["Author & Year", "Methodology", "Key Findings", "Limitations"]
CW = [1.75, 2.25, 2.15, 2.15]
LIT = [
    ("Mangrove & habitat mapping", [
        [[[("Ghorbanian et al., 2025", {"bold": True})], [("IEEE JSTARS ★ foundation paper", {"size": 9, "italic": True})]],
         "Weakly supervised U-Net + EfficientNet-B7 (UNB7) on time-series Sentinel-1 radar (VV/VH).",
         "About 95.56% overall accuracy on their 2020 data; clean maps even from imperfect labels.",
         "Gives a habitat map only — no patch or network analysis; sparse mangroves harder."],
        [[[("Bunting et al., 2018", {"bold": True})], [("Remote Sensing (Global Mangrove Watch)", {"size": 9, "italic": True})]],
         "Classified ALOS radar + Landsat images to map mangrove extent worldwide.",
         "First consistent global mangrove baseline map; widely used by researchers and agencies.",
         "Shows where mangroves are, not how important each area is for connectivity."]]),
    ("Deep-learning segmentation models", [
        [[[("Ronneberger et al., 2015", {"bold": True})], [("MICCAI (U-Net)", {"size": 9, "italic": True})]],
         "Encoder–decoder network with skip connections that labels every pixel.",
         "Accurate pixel-level segmentation even with little training data; now a standard model.",
         "Built for medical images; needs labelled masks; no idea of habitat networks."],
        [[[("Tan & Le, 2019", {"bold": True})], [("ICML (EfficientNet)", {"size": 9, "italic": True})]],
         "Scales network depth, width and image size together (models B0 to B7).",
         "Higher accuracy with fewer parameters than earlier image-recognition networks.",
         "Large versions (B7) need a lot of GPU memory; it is a backbone, not a full mapping system."]]),
    ("Transformers & change detection", [
        [[[("Liu et al., 2021", {"bold": True})], [("ICCV (Swin Transformer)", {"size": 9, "italic": True})]],
         "Vision transformer that looks at the image in shifted windows at several scales.",
         "Strong results in classification, detection and segmentation benchmarks.",
         "Needs heavy compute and large data; general vision model, not habitat-specific."],
        [[[("Yu et al., 2025", {"bold": True})], [("IEEE JSTARS (review)", {"size": 9, "italic": True})]],
         "Survey of deep-learning change detection: CNN, Siamese and transformer methods.",
         "Organises methods and datasets for finding change between two satellite images.",
         "Survey only; finds change but does not measure its effect on habitat connectivity."]]),
    ("Graph-based landscape connectivity", [
        [[[("Urban & Keitt, 2001", {"bold": True})], [("Ecology", {"size": 9, "italic": True})]],
         "Represents a landscape as a graph: habitat patches as nodes, movement links as edges.",
         "Graph measures show which patches and links keep a landscape connected.",
         "Patches entered by hand; no satellite images or automation."],
        [[[("McRae et al., 2008", {"bold": True})], [("Ecology (circuit theory)", {"size": 9, "italic": True})]],
         "Treats the landscape like an electrical circuit to model all possible movement paths.",
         "More realistic corridor modelling than a single shortest path.",
         "Computationally heavy for large areas; needs a prepared resistance map."]]),
    ("Connectivity indices", [
        [[[("Pascual-Hortal & Saura, 2006", {"bold": True})], [("Landscape Ecology", {"size": 9, "italic": True})]],
         "Compared graph connectivity indices and proposed the Integral Index of Connectivity (IIC).",
         "IIC can rank patches by how much connectivity is lost when each is removed.",
         "Needs habitat patches already mapped; not linked to satellite AI."],
        [[[("Saura & Pascual-Hortal, 2007", {"bold": True})], [("Landscape & Urban Planning", {"size": 9, "italic": True})]],
         "Probability of Connectivity (PC) index using the chance of movement between patches.",
         "Basis of the Conefor tool, widely used in conservation planning.",
         "Needs dispersal distances per species and pre-made habitat maps."]]),
    ("Graph learning on networks", [
        [[[("Liu et al., 2022", {"bold": True})], [("IEEE TGRS", {"size": 9, "italic": True})]],
         "Spatio-temporal graph neural network on long-term satellite image series.",
         "Detects how landscape patterns change over time.",
         "Finds the type of change only; does not rank which patches are critical."],
        [[[("Liu & Liu, 2025", {"bold": True})], [("IEEE Access (review)", {"size": 9, "italic": True})]],
         "Systematic review of graph neural networks for infrastructure reliability and resilience.",
         "Graph models can estimate node importance and how a network survives failures.",
         "Engineering networks only — not ecology or satellite habitat maps."]]),
    ("Explainable AI in remote sensing", [
        [[[("Naceur et al., 2025", {"bold": True})], [("IEEE TGRS", {"size": 9, "italic": True})]],
         "U-Net on spectral indices with SHAP explanations for change detection.",
         "Robust change detection that explains which image bands drove the change.",
         "Explains pixels and bands, not the role of a habitat patch in a network."],
        [[[("Klotz et al., 2025", {"bold": True})], [("IEEE JSTARS", {"size": 9, "italic": True})]],
         "Benchmark of explanation methods (Occlusion, LIME, Grad-CAM, LRP, DeepLIFT).",
         "How useful an explanation method is varies a lot by method and metric.",
         "Explains image-level predictions, not network-level decisions."]]),
]
for k, (theme, rows) in enumerate(LIT, 1):
    s = new_slide(f"Literature Survey ({k}/7)", size=28)
    text(s, 1.3, 1.1, 8.3, 0.35, theme, size=14, bold=True, color=NAVY)
    table(s, 1.3, 1.5, 8.3, CW, rows, HDR, size=12.5, head_size=13, row_h=1.6)
    if k == 7:
        banner(s, 1.3, 5.2, 8.3, 0.9,
               [("Research gap:  ", {"bold": True, "color": GOLD}),
                ("these works map habitat, model connectivity or explain pixels separately. None turn satellite "
                 "predictions into an explained network of patches for what-if and restoration decisions — "
                 "that is what EcoConnectAI adds.", {})], size=12)
    notes(s, f"Literature survey part {k} of 7 — {theme}. For each paper: what they did, what they found, and what "
             "is missing that our project addresses.")

# =============================================================================== 10 · PROBLEM + OBJECTIVES
s = new_slide("Problem Statement and Objectives", size=28)
box(s, 1.3, 1.2, 4.0, 3.2, fill=PEACH)
badge(s, 1.5, 1.37, 0.42, "!", ORANGE, size=16)
text(s, 2.05, 1.38, 3.1, 0.45, "Problem Statement", size=17, bold=True, color=ORANGE)
yy = 2.0
for p in ["Maps show where mangroves are — not which patches matter.",
          "Two patches can look alike, but one may be the only link between others.",
          "If that small link is lost, the whole network can break apart."]:
    box(s, 1.55, yy + 0.1, 0.1, 0.1, fill=ORANGE, shape=MSO_SHAPE.OVAL)
    text(s, 1.75, yy, 3.4, 0.7, p, size=12.5)
    yy += 0.75
box(s, 5.55, 1.2, 4.05, 3.2, fill=MINT)
badge(s, 5.75, 1.37, 0.42, "✓", GREEN, size=15)
text(s, 6.3, 1.38, 3.1, 0.45, "Objectives", size=17, bold=True, color=GREEN)
yy = 1.95
for o in ["Detect mangrove habitat from satellite images", "Divide the habitat into patches",
          "Find the important connecting patches", "Support protection & restoration decisions"]:
    lbox(s, 5.8, yy, 3.55, 0.48, o, fill=WHITE, line=GREEN, color=NAVY, size=12)
    yy += 0.58
text(s, 1.3, 4.6, 1.9, 0.4, "Existing approach", size=12.5, bold=True, color=MUTED)
lbox(s, 3.2, 4.55, 1.7, 0.45, "Satellite map", fill=LIGHT, line=LINE, color=TEXT, size=11.5)
arrow(s, 4.97, 4.66)
lbox(s, 5.32, 4.55, 1.9, 0.45, "Where is habitat?", fill=LIGHT, line=LINE, color=TEXT, size=11.5)
text(s, 1.3, 5.25, 1.9, 0.4, "Our approach", size=12.5, bold=True, color=NAVY)
for i, (st, x) in enumerate(zip(["Satellite map", "Where is habitat?", "How is it connected?", "Which patches matter?"],
                                [3.2, 4.73, 6.26, 7.79])):
    last = i == 3
    lbox(s, x, 5.2, 1.8 if last else 1.3, 0.55, st, fill=NAVY if last else WHITE, color=WHITE if last else NAVY, size=11)
    if not last:
        arrow(s, x + 1.33, 5.37, w=0.2, h=0.2)
notes(s, "The problem: finding where mangroves are is not enough. Our objectives: detect the habitat, split it into "
         "patches, find the important connecting patches, and support protection and restoration decisions.")

# =============================================================================== 11–12 · EXISTING SYSTEM
s = new_slide("Existing System — Architecture", size=28)
text(s, 1.3, 1.15, 8.3, 0.45, "How mangrove habitat is studied today (Google Earth Engine, ArcGIS / QGIS, Global Mangrove Watch)",
     size=12.5, color=MUTED, italic=True)
steps = [("Satellite images", "Landsat / Sentinel"), ("GIS software", "Earth Engine, ArcGIS, QGIS"),
         ("Manual analysis", "by a trained GIS expert"), ("Static habitat map", "e.g. Global Mangrove Watch"),
         ("Report", "maps and area statistics")]
x = 1.3
for i, (h, d) in enumerate(steps):
    lbox(s, x, 1.85, 1.46, 1.05, h, fill=LIGHT, line=MUTED, color=TEXT, size=12, sub=d, subsize=9.5)
    if i < 4:
        arrow(s, x + 1.49, 2.26, w=0.2, h=0.2, color=MUTED)
    x += 1.71
box(s, 1.3, 3.25, 8.3, 1.3, fill=WHITE, line=LINE)
text(s, 1.5, 3.33, 8.0, 0.4, "What the existing system answers", size=14, bold=True, color=NAVY)
text(s, 1.5, 3.75, 8.0, 0.75, [[("✓  ", {"color": GREEN, "bold": True}), ("Where is the mangrove?  How much area is there?", {})],
                               [("✗  ", {"color": ORANGE, "bold": True}), ("Which patch keeps the network connected?  What if it is lost?  Where to restore first?", {})]],
     size=12.5, space=4)
banner(s, 1.3, 4.85, 8.3, 0.75, [("Output is a static map. ", {"bold": True, "color": GOLD}),
                                  ("Each habitat patch is shown on its own, with no information about how patches are connected.", {})],
       size=12.5)
notes(s, "Today, habitat is studied with tools like Google Earth Engine, ArcGIS or QGIS: an expert processes the images "
         "and produces a static habitat map such as Global Mangrove Watch. This answers where the mangroves are, but not "
         "which patches matter.")

s = new_slide("Existing System — Drawbacks and Limitations", size=26)
draw = [("Shows location, not importance", "A map says where habitat is, not which patch holds the network together."),
        ("No what-if analysis", "Cannot test what happens if a patch is lost to a road, port or cyclone."),
        ("No restoration ranking", "Does not suggest which area to restore first."),
        ("Needs GIS experts", "Hard for forest officers and policymakers to use directly."),
        ("Optical images blocked by cloud", "Monsoon cloud hides the coast for weeks in optical satellite images."),
        ("No evidence or follow-up", "No explanation of results and no link to field verification.")]
for i, (h, d) in enumerate(draw):
    col, row = i % 2, i // 2
    x, y = 1.3 + col * 4.2, 1.25 + row * 1.5
    box(s, x, y, 4.05, 1.35, fill=PEACH)
    badge(s, x + 0.15, y + 0.15, 0.38, "!", ORANGE, size=13)
    text(s, x + 0.65, y + 0.12, 3.3, 1.15, [[(h, {"bold": True, "color": ORANGE, "size": 13})], [(d, {"size": 11})]], space=2)
notes(s, "The main drawbacks: existing tools show location but not importance, cannot run what-if tests, do not rank "
         "restoration sites, need GIS experts, depend on cloud-free optical images, and give no explanation or link to "
         "field verification.")

# =============================================================================== 13–16 · PROPOSED SYSTEM
s = new_slide("Proposed System — Architecture", size=28)
text(s, 1.3, 1.1, 5, 0.35, "How the analysis works", size=14, bold=True, color=NAVY)
W, H = 2.35, 0.6
cols, rows = [1.3, 4.25, 7.2], [1.5, 2.5, 3.5]
r1 = ["Sentinel-1 satellite data", "Preprocessing", "AI mangrove detection"]
r2 = ["Mangrove probability map", "Habitat patch extraction", "Connectivity network"]
r3 = ["Critical patches + what-if", "Restoration priority", "Web dashboard"]
for i, t in enumerate(r1):
    lbox(s, cols[i], rows[0], W, H, t, size=12.5)
    if i < 2: arrow(s, cols[i] + W + 0.18, rows[0] + 0.19, w=0.26)
arrow(s, cols[2] + W / 2 - 0.12, rows[0] + H + 0.08, w=0.24, h=0.24, direction="down")
for i, t in enumerate(r2):
    lbox(s, cols[2 - i], rows[1], W, H, t, size=12.5)
    if i < 2: arrow(s, cols[2 - i] - 0.44, rows[1] + 0.19, w=0.26, direction="left")
arrow(s, cols[0] + W / 2 - 0.12, rows[1] + H + 0.08, w=0.24, h=0.24, direction="down")
for i, t in enumerate(r3):
    last = i == 2
    lbox(s, cols[i], rows[2], W, H, t, size=12.5, fill=NAVY if last else WHITE, color=WHITE if last else NAVY)
    if i < 2: arrow(s, cols[i] + W + 0.18, rows[2] + 0.19, w=0.26)
text(s, 1.3, 4.4, 5, 0.35, "Software that runs it", size=14, bold=True, color=GREEN)
x = 1.3
for i, (h, d) in enumerate([("Python + PyTorch", "AI model"), ("FastAPI", "backend server"), ("SQLite", "database"),
                            ("Next.js + React", "web frontend"), ("Leaflet + React Flow", "map + network view")]):
    lbox(s, x, 4.8, 1.5, 0.8, h, fill=MINT, line=GREEN, color=GREEN, size=11.5, sub=d, subsize=9.5)
    if i < 4: arrow(s, x + 1.53, 5.1, w=0.14, h=0.18, color=GREEN)
    x += 1.7
notes(s, "Top: the analysis flow from satellite data to the web dashboard. Bottom: the software — Python and PyTorch for "
         "the AI, a FastAPI backend with a SQLite database, and a Next.js/React frontend with a Leaflet map and a React "
         "Flow network view.")

s = new_slide("Proposed System — Advantages", size=28)
adv = [("Shows which patches matter", "Ranks every patch by how much connectivity is lost without it."),
       ("Instant what-if tests", "Remove or restore a patch and see the new result straight away."),
       ("Restoration ranking", "Candidate areas ranked by the connectivity they would add."),
       ("Works through cloud", "Radar images (Sentinel-1) see through monsoon cloud."),
       ("Explains every result", "Plain-language reason with the numbers behind each ranking."),
       ("Easy to use", "One web dashboard with roles, alerts, field tasks and reports.")]
for i, (h, d) in enumerate(adv):
    col, row = i % 2, i // 2
    x, y = 1.3 + col * 4.2, 1.25 + row * 1.5
    box(s, x, y, 4.05, 1.35, fill=MINT)
    badge(s, x + 0.15, y + 0.15, 0.38, "✓", GREEN, size=12)
    text(s, x + 0.65, y + 0.12, 3.3, 1.15, [[(h, {"bold": True, "color": GREEN, "size": 13})], [(d, {"size": 11})]], space=2)
notes(s, "Compared with the existing system, ours shows which patches matter, runs what-if tests instantly, ranks "
         "restoration areas, works through cloud, explains every result, and is usable through one dashboard.")

s = new_slide("Proposed System — Module Explanation", size=26)
mods = [("Data acquisition", "Downloads free Sentinel-1 radar scenes for each study area and cuts them into tiles."),
        ("Mangrove detection", "A deep-learning model marks mangrove pixels and gives a probability map."),
        ("Patch extraction", "Touching mangrove pixels are grouped into habitat patches (at least 2 ha)."),
        ("Connectivity network", "Each patch is linked to its nearest patches within 5 km."),
        ("Criticality & what-if", "Removes patches one at a time, or on request, and measures the change."),
        ("Restoration & dashboard", "Ranks candidate areas; map, network view, alerts, field tasks and reports.")]
for i, (h, d) in enumerate(mods, 1):
    col, row = (i - 1) % 2, (i - 1) // 2
    x, y = 1.3 + col * 4.2, 1.25 + row * 1.5
    box(s, x, y, 4.05, 1.35, fill=LIGHT)
    badge(s, x + 0.15, y + 0.15, 0.4, str(i), NAVY, size=13)
    text(s, x + 0.68, y + 0.12, 3.3, 1.15, [[(h, {"bold": True, "color": NAVY, "size": 13})], [(d, {"size": 11})]], space=2)
notes(s, "The system has six modules: data acquisition, mangrove detection, patch extraction, the connectivity network, "
         "criticality and what-if analysis, and restoration with the dashboard.")

s = new_slide("Proposed System — Algorithms Used", size=28)
algs = [
    ["U-Net with EfficientNet-B0", "Mangrove detection", "A deep-learning model that labels each 10 m pixel as mangrove or not."],
    ["Connected-component labelling", "Patch extraction", "Groups touching mangrove pixels into one patch; removes tiny pieces."],
    ["Nearest-neighbour network (5 km)", "Connectivity network", "Links each patch to its 3 nearest patches if they are within 5 km."],
    ["Integral Index of Connectivity (IIC)", "Connectivity score", "One number for the network: large patches reachable in few steps score high."],
    ["Remove-one-patch test", "Criticality", "Delete a patch, recompute the score; the bigger the drop, the more critical."],
    ["Add-one-candidate test", "Restoration priority", "Add a candidate area, recompute; rank areas by the gain."]]
table(s, 1.3, 1.25, 8.3, [2.55, 1.85, 3.9], algs, ["Algorithm", "Used for", "In simple words"], size=12.5, head_size=13, row_h=0.66)
text(s, 1.3, 5.75, 8.3, 0.45, "Tested at 3, 5 and 8 km to check the 5 km choice; the model threshold was tuned on test tiles.",
     size=10.5, color=MUTED, italic=True)
notes(s, "The main algorithms: a U-Net deep-learning model for mangrove detection, connected-component labelling to form "
         "patches, a nearest-neighbour network within 5 km, the Integral Index of Connectivity as the network score, a "
         "remove-one-patch test for criticality, and an add-one-candidate test for restoration.")

# =============================================================================== 17–20 · RESULTS AND DISCUSSION
s = new_slide("Results — Mangrove Detection", size=28)
for x, val, lab, sub in [(1.3, "0.842", "Test IoU", "overlap with reference map"), (3.45, "0.914", "Test F1 score", "precision + recall")]:
    box(s, x, 1.2, 2.0, 1.45, fill=WHITE, line=NAVY, lw=1.75)
    text(s, x, 1.25, 2.0, 0.72, val, size=32, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    text(s, x, 1.95, 2.0, 0.65, [[(lab, {"bold": True, "size": 12.5})], [(sub, {"size": 10, "color": MUTED})]],
         align=PP_ALIGN.CENTER, space=0)
box(s, 5.65, 1.2, 3.95, 1.45, fill=PEACH)
text(s, 5.8, 1.25, 3.7, 1.38, anchor=MSO_ANCHOR.MIDDLE, paras=[
    [("Development model result", {"bold": True, "color": ORANGE, "size": 13})],
    [("200 test tiles, measured against Global Mangrove Watch reference labels — not field accuracy.", {"size": 11})]], space=2)
text(s, 1.3, 2.85, 5.2, 0.35, "Comparison: test IoU of our development models", size=13, bold=True, color=NAVY)
cd = CategoryChartData()
cd.categories = ["4 areas · radar", "Kerala · radar", "Kerala · optical", "Kerala · radar+optical"]
cd.add_series("Test IoU", (0.842, 0.023, 0.054, 0.053))
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(1.3), Inches(3.2), Inches(5.1), Inches(2.75), cd)
ch = gf.chart; ch.has_legend = False; ch.has_title = False
ser = ch.series[0]; ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(NAVY)
ser.data_labels.show_value = True; ser.data_labels.number_format = "0.000"; ser.data_labels.number_format_is_linked = False
ser.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END; ser.data_labels.font.size = Pt(10); ser.data_labels.font.bold = True
ch.value_axis.visible = False; ch.value_axis.has_major_gridlines = False; ch.value_axis.maximum_scale = 1.0
ch.value_axis.minimum_scale = 0
ch.category_axis.tick_labels.font.size = Pt(10); ch.category_axis.reverse_order = True
ch.category_axis.format.line.color.rgb = rgb(LINE)
box(s, 6.6, 2.85, 3.0, 3.1, fill=LIGHT)
text(s, 6.75, 2.95, 2.75, 2.95, [
    [("What this means", {"bold": True, "color": NAVY, "size": 13})],
    [("Training on all four areas works far better than Kerala alone.", {"size": 11})],
    [("Most of the score comes from the Sundarbans; thin Kerala mangrove strips are still hard.", {"size": 11})],
    [("95.56% is the published foundation-paper result, not ours.", {"size": 11, "italic": True, "color": MUTED})]], space=5)
notes(s, "Our development model scores test IoU 0.842 and F1 0.914 on 200 test tiles, measured against Global Mangrove "
         "Watch reference labels. The comparison shows that training on four areas is much better than Kerala only. "
         "Most of the strength comes from the Sundarbans; thin Kerala mangroves remain difficult.")

s = new_slide("Results — System Outputs", size=28)
picture(s, "dashboard_map.png", 1.3, 1.2, w=4.1)
text(s, 1.3, 3.72, 4.1, 0.35, "Dashboard map: patches, links, critical patches, candidates", size=10, color=MUTED,
     italic=True, align=PP_ALIGN.CENTER)
picture(s, "graph.png", 5.55, 1.2, w=4.05)
text(s, 5.55, 3.72, 4.05, 0.35, "Network view: 24 patches, 43 links, 2 groups (Kerala)", size=10, color=MUTED,
     italic=True, align=PP_ALIGN.CENTER)
kp = [("24", "habitat patches"), ("43", "connectivity links"), ("≈ 220 ha", "mangrove mapped"), ("4", "restoration candidates")]
x = 1.3
for v, l in kp:
    box(s, x, 4.2, 1.95, 1.05, fill=LIGHT)
    text(s, x, 4.25, 1.95, 0.55, v, size=22, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    text(s, x, 4.8, 1.95, 0.4, l, size=11, color=MUTED, align=PP_ALIGN.CENTER)
    x += 2.117
text(s, 1.3, 5.45, 8.3, 0.45, "Kerala (Vembanad–Kol) latest analysis run — development model, not final.", size=10.5,
     color=MUTED, italic=True, align=PP_ALIGN.CENTER)
notes(s, "These are real screens from our application for the Kerala study area: the dashboard map with patches and "
         "links, and the network view. The run found 24 patches, 43 links, about 220 hectares of mangrove and 4 "
         "restoration candidates.")

s = new_slide("Results — Which Patch Matters?", size=28)
text(s, 1.3, 1.1, 8.3, 0.4, [[("Size alone does not tell us which patch is important", {"bold": True, "color": NAVY, "size": 16})]])
for x, fill, line, title, rws, col in [
        (1.3, LIGHT, LINE, "P01 — largest patch", [("35.1 ha", "#1 by size"), ("−30.7%", "connectivity if removed"),
                                                  ("No split", "other routes still connect")], NAVY),
        (3.85, PEACH, ORANGE, "P17 — small bridge patch", [("3.1 ha", "only #17 by size"), ("−27.0%", "connectivity if removed"),
                                                           ("2 → 3 groups", "network breaks apart")], ORANGE)]:
    box(s, x, 1.6, 2.4, 2.4, fill=fill, line=line)
    text(s, x + 0.1, 1.67, 2.2, 0.4, title, size=12.5, bold=True, color=col, align=PP_ALIGN.CENTER)
    yy = 2.12
    for big, small in rws:
        text(s, x + 0.1, yy, 2.2, 0.58, [[(big, {"bold": True, "size": 16, "color": col})],
                                         [(small, {"size": 10, "color": MUTED})]], align=PP_ALIGN.CENTER, space=0)
        yy += 0.6
picture(s, "graph_zoom.png", 7.12, 1.6, h=2.4)
text(s, 6.45, 4.02, 3.15, 0.35, "P17 and P07 are the only “bridge” patches", size=10, color=MUTED, italic=True,
     align=PP_ALIGN.CENTER)
box(s, 1.3, 4.5, 8.3, 1.5, fill=WHITE, line=LINE)
text(s, 1.45, 4.55, 8.0, 1.4, [
    [("Discussion", {"bold": True, "color": NAVY, "size": 13})],
    [("• Two of the top-3 critical patches are small (P07 8.8 ha, P17 3.1 ha) — position in the network matters, not just size.", {"size": 11})],
    [("• In Odisha, a 55.6 ha patch ranks above a 910 ha patch for the same reason.", {"size": 11})],
    [("• The ranking stayed stable when the link distance was changed to 3 km or 8 km (correlation ≥ 0.96 in Kerala).", {"size": 11})]], space=2)
notes(s, "The key finding: P01 is the largest patch, but P17 is only 3.1 hectares and 17th by size — yet removing it cuts "
         "connectivity by 27% and splits the network into three groups. The same pattern appears in Odisha. And the "
         "ranking stays stable when we change the 5 km link distance to 3 or 8 km.")

s = new_slide("Results — What-if and Restoration", size=28)
box(s, 1.3, 1.15, 4.1, 4.85, fill=WHITE, line=LINE)
text(s, 1.45, 1.22, 3.9, 0.45, "What happens if a patch is lost?", size=14.5, bold=True, color=NAVY)
cd = CategoryChartData()
cd.categories = ["Remove P17 (3.1 ha)", "Remove P01 (35.1 ha)"]
cd.add_series("Habitat lost (%)", (1.4, 16.0))
cd.add_series("Connectivity lost (%)", (27.0, 30.7))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1.4), Inches(1.7), Inches(3.9), Inches(3.1), cd)
ch = gf.chart; ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
ch.legend.font.size = Pt(10)
for ser, colr in zip(ch.series, [MUTED, ORANGE]):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(colr)
    ser.data_labels.show_value = True; ser.data_labels.number_format = '0.0"%"'; ser.data_labels.number_format_is_linked = False
    ser.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END; ser.data_labels.font.size = Pt(10); ser.data_labels.font.bold = True
ch.value_axis.visible = False; ch.value_axis.has_major_gridlines = False; ch.value_axis.maximum_scale = 36
ch.category_axis.tick_labels.font.size = Pt(10); ch.category_axis.format.line.color.rgb = rgb(LINE)
text(s, 1.45, 4.85, 3.9, 1.1, [[("P17 is tiny but costs almost as much connectivity as the biggest patch.", {"size": 11})],
                               [("Simulated under project assumptions.", {"size": 10, "italic": True, "color": MUTED})]], space=1)
box(s, 5.6, 1.15, 4.0, 4.85, fill=WHITE, line=LINE)
text(s, 5.75, 1.22, 3.8, 0.45, "What happens if we restore an area?", size=14.5, bold=True, color=GREEN)
yy = 1.75
for i, t in enumerate(["Candidate C1 (1.6 ha)", "Added to the network", "+1.29% connectivity · 3 new links"]):
    last = i == 2
    lbox(s, 5.85, yy, 3.5, 0.42, t, fill=GREEN if last else MINT, line=GREEN, color=WHITE if last else GREEN, size=11.5)
    if not last:
        arrow(s, 7.48, yy + 0.44, w=0.22, h=0.14, direction="down", color=GREEN)
    yy += 0.6
picture(s, "restoration.png", 5.75, 3.6, w=3.7)
text(s, 5.75, 5.1, 3.75, 0.85, [[("Candidates come from model outputs — not field-confirmed.", {"size": 10})],
                                [("No cost data yet, so ranking is by connectivity gain.", {"size": 10, "italic": True,
                                                                                            "color": MUTED})]], space=0)
notes(s, "Left: removing P17 loses only 1.4% of habitat but 27% of connectivity; removing P01 loses 16% of habitat and "
         "30.7% of connectivity. These are simulations. Right: adding candidate area C1 improves connectivity by 1.29% "
         "with three new links. Candidates are not field-confirmed and there is no cost data yet.")

# =============================================================================== 21 · CONCLUSION
s = new_slide("Conclusion", size=30)
box(s, 1.3, 1.2, 8.3, 1.2, fill=LIGHT)
text(s, 1.5, 1.27, 7.9, 1.1, [[("Summary", {"bold": True, "color": NAVY, "size": 14})],
                              [("EcoConnectAI combines satellite radar images, AI mangrove detection, a network of habitat patches "
                                "and what-if analysis to help decide which coastal habitat to protect or restore first.", {"size": 12})]], space=2)
box(s, 1.3, 2.6, 4.05, 2.65, fill=MINT)
text(s, 1.45, 2.67, 3.8, 0.4, "Key findings", size=14, bold=True, color=GREEN)
bullets(s, 1.45, 3.12, 3.8, ["Working end-to-end system on real data for 4 coastal areas",
                             "Development model: IoU 0.842 vs reference labels",
                             "Small patches can be critical: P17 (3.1 ha) → −27% connectivity"], dot=GREEN, size=11.5, gap=0.68, h=0.65)
box(s, 5.55, 2.6, 4.05, 2.65, fill=PEACH)
text(s, 5.7, 2.67, 3.8, 0.4, "Future scope", size=14, bold=True, color=ORANGE)
bullets(s, 5.7, 3.12, 3.8, ["Train the full UNB7 model on a GPU", "Field validation of results",
                            "Real restoration cost data", "Larger-scale deployment for forest departments"],
        dot=ORANGE, size=11.5, gap=0.5, h=0.5)
banner(s, 1.3, 5.5, 8.3, 0.5, [("From mapping habitat to understanding which habitat matters.",
                                 {"bold": True, "italic": True, "color": GOLD})], size=14)
notes(s, "To summarise, EcoConnectAI goes from satellite images to a decision-support network. Our key findings are the "
         "working system, the development model result and the fact that small patches can be critical. Next: train the "
         "full UNB7 model, validate in the field, add cost data and scale up.")

# =============================================================================== 22 · CO-GUIDE MEETING PROOF
s = new_slide("Co-guide Meeting Proof", size=30)
text(s, 1.3, 1.1, 8.3, 0.4, "Co-guide: [Name, Designation]   ·   Guide: Dr. Anusha Preetham (Associate Professor)",
     size=12.5, color=NAVY, bold=True)
mt = [["[dd-mm-2026]", "[Topic discussed / feedback received]", ""],
      ["[dd-mm-2026]", "[Topic discussed / feedback received]", ""],
      ["[dd-mm-2026]", "[Topic discussed / feedback received]", ""]]
table(s, 1.3, 1.55, 8.3, [1.6, 4.9, 1.8], mt, ["Date", "Discussion", "Signature"], size=11, row_h=0.45)
for x in (1.3, 5.55):
    b = box(s, x, 3.6, 4.05, 2.35, fill=None, line=MUTED, dash=True)
    text(s, x, 4.45, 4.05, 0.6, "Insert meeting photo / screenshot here", size=12, color=MUTED, italic=True,
         align=PP_ALIGN.CENTER)
notes(s, "Replace the bracketed fields with your co-guide's name, meeting dates and discussion points, and insert the "
         "meeting photos or signed log.")

# =============================================================================== 23 · SDG GOAL
s = new_slide("SDG Goal", size=30)
text(s, 1.3, 1.1, 8.3, 0.4, "EcoConnectAI maps to three of the 17 UN Sustainable Development Goals", size=13, color=MUTED,
     italic=True)
sdg = [("14", "Life Below Water", "0A97D9",
        "Mangroves are nurseries for fish and protect coastal marine life. We help find and protect the most important patches."),
       ("13", "Climate Action", "3F7E44",
        "Mangroves store carbon and shield coasts from storms. Keeping them connected keeps these services working."),
       ("15", "Life on Land", "56C02B",
        "Stopping habitat fragmentation and guiding restoration of degraded ecosystems.")]
x = 1.3
for num, name, colr, why in sdg:
    box(s, x, 1.65, 2.62, 3.95, fill=WHITE, line=colr, lw=2)
    sq = box(s, x + 0.61, 1.85, 1.4, 1.4, fill=colr, radius=0.1)
    tf = sq.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = num; r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = rgb(WHITE); r.font.name = FONT
    text(s, x + 0.1, 3.35, 2.42, 0.45, f"SDG {num}: {name}", size=13.5, bold=True, color=colr, align=PP_ALIGN.CENTER)
    text(s, x + 0.15, 3.85, 2.32, 1.7, why, size=11, align=PP_ALIGN.CENTER)
    x += 2.84
notes(s, "Our project maps mainly to SDG 14, Life Below Water, because mangroves support coastal marine life; SDG 13, "
         "Climate Action, because mangroves store carbon and protect coasts; and SDG 15, Life on Land, because we address "
         "habitat fragmentation and restoration.")

# =============================================================================== 24 · STAKEHOLDERS (college template)
def fill_bullets(shape, items, size=10):
    """items: list of (label, value) — reuse the template's bullet paragraph formatting."""
    txBody = shape.text_frame._txBody
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    tmpl = TEMPLATE_P
    for p in txBody.findall(A + "p"):
        txBody.remove(p)
    for label, value in items:
        p = copy.deepcopy(tmpl)
        runs = p.findall(A + "r")
        if label:
            runs[0].find(A + "t").text = label; runs[1].find(A + "t").text = " " + value
        else:
            p.remove(runs[0]); runs[1].find(A + "t").text = value
        for r in p.findall(A + "r"):
            r.find(A + "rPr").set("sz", str(int(size * 100)))
        txBody.append(p)


by_name = {sh.name: sh for sh in stake_slide.shapes}
TEMPLATE_P = copy.deepcopy(by_name["Text 15"].text_frame._txBody.findall(
    "{http://schemas.openxmlformats.org/drawingml/2006/main}p")[0])
fill_bullets(by_name["Text 8"], [("Problem source:", "research gap found in the foundation paper"),
                                 ("Intended client:", "coastal forest & wetland departments")])
fill_bullets(by_name["Text 15"], [("Target Users:", "forest range officers, GIS analysts, researchers"),
                                  ("User Need:", "know which patches matter and test protect / restore options")])
fill_bullets(by_name["Text 22"], [("", "Coastal communities (fishing, storm protection)"),
                                  ("", "Conservation NGOs and ecologists"),
                                  ("", "Policymakers and coastal planning authorities")])
fill_bullets(by_name["Text 29"], [("Existing:", "GMW / Earth Engine maps show where it is"),
                                  ("Our Project:", "ranks patches; what-if and restoration"),
                                  ("Key Improvement:", "“which habitat matters?”")])
kd = by_name["Text 31"].text_frame.paragraphs[0].runs
kd[1].text = ("EcoConnectAI turns a satellite habitat map into a network and measures how much connectivity is lost when "
              "each patch is removed — so a small but critical patch is found, explained and testable before it is lost.")
for r in kd:
    r.font.size = Pt(10)
stake_slide.notes_slide.notes_text_frame.text = (
    "Client: the problem comes from the research gap; the intended client is coastal forest and wetland departments — "
    "they are not sponsors. End users are range officers, GIS analysts and researchers. Other stakeholders are coastal "
    "communities, NGOs and policymakers. Our difference: existing tools show where mangroves are; ours shows which "
    "patches matter and lets you test loss and restoration.")
order.append(stake_slide)

# =============================================================================== 25 · REFERENCES
s = new_slide("References", size=30)
refs = [
    "A. Ghorbanian, A. Ghorbanian, S. A. Ahmadi, A. Mohammadzadeh, and A. Naboureh, “Weakly supervised semantic segmentation of mangrove ecosystem using Sentinel-1 SAR and deep convolutional neural networks,” IEEE J. Sel. Topics Appl. Earth Observ. Remote Sens., vol. 18, pp. 17497–17512, 2025.",
    "P. Bunting et al., “The Global Mangrove Watch — a new 2010 global baseline of mangrove extent,” Remote Sensing, vol. 10, no. 10, p. 1669, 2018.",
    "O. Ronneberger, P. Fischer, and T. Brox, “U-Net: Convolutional networks for biomedical image segmentation,” in Proc. MICCAI, 2015, pp. 234–241.",
    "M. Tan and Q. Le, “EfficientNet: Rethinking model scaling for convolutional neural networks,” in Proc. ICML, 2019, pp. 6105–6114.",
    "Z. Liu et al., “Swin transformer: Hierarchical vision transformer using shifted windows,” in Proc. IEEE/CVF ICCV, 2021, pp. 10012–10022.",
    "C. Yu et al., “Deep learning-based change detection in remote sensing: A comprehensive review,” IEEE J. Sel. Topics Appl. Earth Observ. Remote Sens., vol. 18, pp. 24415–24437, 2025.",
    "D. Urban and T. Keitt, “Landscape connectivity: A graph-theoretic perspective,” Ecology, vol. 82, no. 5, pp. 1205–1218, 2001.",
    "B. H. McRae, B. G. Dickson, T. H. Keitt, and V. B. Shah, “Using circuit theory to model connectivity in ecology, evolution, and conservation,” Ecology, vol. 89, no. 10, pp. 2712–2724, 2008.",
    "L. Pascual-Hortal and S. Saura, “Comparison and development of new graph-based landscape connectivity indices,” Landscape Ecology, vol. 21, no. 7, pp. 959–967, 2006.",
    "S. Saura and L. Pascual-Hortal, “A new habitat availability index to integrate connectivity in landscape conservation planning,” Landscape and Urban Planning, vol. 83, no. 2–3, pp. 91–103, 2007.",
    "M. Liu et al., “Hybrid spatiotemporal graph convolutional network for detecting landscape pattern evolution from long-term remote sensing images,” IEEE Trans. Geosci. Remote Sens., vol. 60, pp. 1–16, 2022.",
    "T. Liu and F. Liu, “Graph neural networks for evaluating the reliability and resilience of infrastructure systems: A systematic review,” IEEE Access, vol. 13, pp. 164883–164904, 2025.",
    "Y. Naceur, S. Bouzidi, and M. Zaouali, “U-Net for remote sensing: A spectral index-based approach with explainable AI for robust change detection,” IEEE Trans. Geosci. Remote Sens., vol. 63, pp. 1–7, 2025.",
    "J. Klotz, T. Burgert, and B. Demir, “On the effectiveness of methods and metrics for explainable AI in remote sensing image scene classification,” IEEE J. Sel. Topics Appl. Earth Observ. Remote Sens., vol. 18, pp. 27764–27780, 2025.",
]
text(s, 1.3, 1.1, 8.35, 5.0, [[(f"[{i}]  ", {"bold": True, "color": NAVY}), (r, {})] for i, r in enumerate(refs, 1)],
     size=9.5, space=2.5)
notes(s, "References in IEEE style.")

# =============================================================================== 26 · THANK YOU
s = new_slide("Thank You", size=48)
s.shapes.title.top = Inches(1.9); s.shapes.title.height = Inches(1.2)
text(s, 1.3, 3.15, 8.3, 0.6, "Queries?", size=26, color=NAVY, align=PP_ALIGN.CENTER)
text(s, 1.3, 4.0, 8.3, 0.9, [[("EcoConnectAI", {"bold": True, "color": NAVY})],
                             [("Kuldeep Raj (1DS23CS112)  ·  Ruhinaaz (1DS23CS187)", {"color": MUTED, "size": 12})]],
     size=15, align=PP_ALIGN.CENTER, space=2)
notes(s, "Thank you. We are happy to take questions.")

# =============================================================================== final order
new_ids = [prs.slides._sldIdLst[prs.slides.index(sl)] for sl in order]
for e in list(prs.slides._sldIdLst):
    prs.slides._sldIdLst.remove(e)
for e in new_ids:
    prs.slides._sldIdLst.append(e)
prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")
