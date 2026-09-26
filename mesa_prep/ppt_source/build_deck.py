"""Build the 9-slide EcoConnectAI final presentation on the college (DSI) template.

Keeps the original slide master/layout (blue strip, yellow DSI footer, logo) from the college deck and
replaces all slides. Every figure comes from the project's run/evaluation files (see notes per slide).
"""
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

SRC = "orig.pptx"
OUT = "EcoConnectAI_Final_Presentation.pptx"
IMG = "img/"

NAVY, ORANGE, PEACH = "0D3889", "D96A2B", "FBEFE6"
GREEN, MINT, LIGHT = "0E9E74", "E7F6F1", "EEF2F8"
TEXT, MUTED, WHITE, LINE = "1A1A1A", "5A6370", "FFFFFF", "C9D3E6"
FONT = "Calibri"


def rgb(h):
    return RGBColor.from_string(h)


prs = Presentation(SRC)
# ---- drop every existing slide, keep master + layout (college theme)
sld = prs.slides._sldIdLst
for s in list(sld):
    prs.part.drop_rel(s.rId)
    sld.remove(s)
LAYOUT = prs.slide_layouts[0]          # "OBJECT" layout used by every slide of the college deck


def new_slide(title, size=32):
    s = prs.slides.add_slide(LAYOUT)
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx != 0:
            ph._element.getparent().remove(ph._element)
    t = s.shapes.title
    t.text_frame.text = title
    p = t.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    for r in p.runs:
        r.font.size = Pt(size); r.font.name = FONT; r.font.color.rgb = rgb(TEXT)
    t.left, t.top, t.width, t.height = Inches(1.15), Inches(0.25), Inches(8.6), Inches(0.95)
    return s


def box(s, x, y, w, h, fill=WHITE, line=None, radius=0.12, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=1.0):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def text(s, x, y, w, h, paras, size=14, color=TEXT, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         italic=False, space=4):
    """paras: str | list of str | list of list[(text, {opts})]"""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04); tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space)
        runs = para if isinstance(para, list) else [(para, {})]
        for rt, o in runs:
            r = p.add_run(); r.text = rt
            f = r.font; f.name = FONT; f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic); f.color.rgb = rgb(o.get("color", color))
    return tb


def badge(s, x, y, d, label, fill, size=14):
    c = box(s, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label; r.font.size = Pt(size); r.font.bold = True; r.font.name = FONT
    r.font.color.rgb = rgb(WHITE)
    return c


def labelled_box(s, x, y, w, h, label, fill=WHITE, line=NAVY, color=NAVY, size=12.5, bold=True, sub=None, subsize=10):
    b = box(s, x, y, w, h, fill=fill, line=line)
    tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.06); tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label; r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT
    r.font.color.rgb = rgb(color)
    if sub:
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run(); r2.text = sub; r2.font.size = Pt(subsize); r2.font.name = FONT
        r2.font.color.rgb = rgb(MUTED if fill == WHITE or fill in (LIGHT, MINT, PEACH) else WHITE)
    return b


def arrow(s, x, y, w=0.28, h=0.22, direction="right", color=NAVY):
    shp = {"right": MSO_SHAPE.RIGHT_ARROW, "down": MSO_SHAPE.DOWN_ARROW, "left": MSO_SHAPE.LEFT_ARROW}[direction]
    return box(s, x, y, w, h, fill=color, shape=shp)


def picture(s, path, x, y, w=None, h=None, border=LINE):
    kw = {}
    if w: kw["width"] = Inches(w)
    if h: kw["height"] = Inches(h)
    pic = s.shapes.add_picture(IMG + path, Inches(x), Inches(y), **kw)
    pic.line.color.rgb = rgb(border); pic.line.width = Pt(1)
    return pic


def banner(s, x, y, w, h, runs, size=13):
    b = box(s, x, y, w, h, fill=NAVY, radius=0.18)
    tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.15)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    for rt, o in runs:
        r = p.add_run(); r.text = rt; r.font.size = Pt(o.get("size", size)); r.font.bold = o.get("bold", False)
        r.font.name = FONT; r.font.italic = o.get("italic", False)
        r.font.color.rgb = rgb(o.get("color", WHITE))
    return b


# =========================================================================== 1 · TITLE
s = new_slide("EcoConnectAI", size=44)
text(s, 1.35, 1.15, 8.1, 0.9,
     [[("A Satellite-Driven Framework for Coastal Ecosystem Connectivity", {})],
      [("and Conservation Decision Support", {})]],
     size=19, color=NAVY, italic=True, align=PP_ALIGN.CENTER, space=0)
text(s, 1.35, 2.35, 8.1, 1.2,
     [[("Presented by", {"size": 12, "color": MUTED})],
      [("Kuldeep Raj  (1DS23CS112)", {})],
      [("Ruhinaaz  (1DS23CS187)", {})]], size=15, align=PP_ALIGN.CENTER, space=2)
text(s, 1.35, 3.75, 8.1, 0.75,
     [[("Guided by", {"size": 12, "color": MUTED})],
      [("Dr. Anusha Preetham  (Associate Professor)", {})]], size=15, align=PP_ALIGN.CENTER, space=2)
text(s, 1.35, 4.75, 8.1, 0.9,
     ["Department of Computer Science and Engineering", "Dayananda Sagar College of Engineering",
      "Final Year Project Evaluation · 2026"], size=12.5, color=MUTED, align=PP_ALIGN.CENTER, space=1)
s.notes_slide.notes_text_frame.text = (
    "Good morning. Our project is EcoConnectAI — a tool that uses satellite images and AI to find mangrove "
    "habitat, and then works out which habitat patches matter most for keeping the coastal ecosystem connected.")

# =========================================================================== 2 · ABSTRACT
s = new_slide("Abstract")
text(s, 1.35, 1.2, 8.2, 0.55,
     "EcoConnectAI uses satellite images and AI to find mangroves, then studies them as a connected network "
     "to see which patches matter most.", size=14, color=MUTED, italic=True)
pts = [("Detect mangroves", "from Sentinel-1 satellite images, pixel by pixel"),
       ("Divide into patches", "group nearby mangrove pixels into habitat patches"),
       ("Connect the patches", "link nearby patches to form a network"),
       ("Find important patches", "remove one patch at a time and see what is lost"),
       ("Support decisions", "test patch loss and restoration on a web dashboard")]
y = 1.95
for i, (h, d) in enumerate(pts, 1):
    badge(s, 1.35, y + 0.05, 0.42, str(i), NAVY, size=13)
    text(s, 1.9, y, 3.3, 0.72, [[(h, {"bold": True, "color": NAVY, "size": 14})], [(d, {"size": 11.5, "color": TEXT})]],
         space=0)
    y += 0.78
picture(s, "dashboard_map.png", 5.35, 1.95, w=4.25)
text(s, 5.35, 4.55, 4.25, 0.5, "Our working dashboard — mangrove patches and links, Vembanad–Kol (Kerala)",
     size=10, color=MUTED, italic=True, align=PP_ALIGN.CENTER)
banner(s, 1.35, 5.95 - 0.1, 8.25, 0.5,
       [("In short:  ", {"bold": True, "color": "FFD54A"}),
        ("from “where are the mangroves?” to “which mangrove patch matters most?”", {})], size=13)
s.notes_slide.notes_text_frame.text = (
    "In simple words: we detect mangroves from satellite images, split them into patches, connect nearby patches "
    "into a network, and remove one patch at a time to see how much connectivity is lost. The dashboard lets an "
    "officer test what happens if a patch is lost, or if an area is restored.")

# =========================================================================== 3 · PROBLEM + OBJECTIVE
s = new_slide("Problem and Objective")
# problem card
box(s, 1.3, 1.25, 4.0, 3.15, fill=PEACH)
badge(s, 1.5, 1.42, 0.42, "!", ORANGE, size=16)
text(s, 2.05, 1.43, 3.1, 0.45, "Problem", size=18, bold=True, color=ORANGE)
probs = ["Maps show where mangroves are — not which patches matter.",
         "Two patches can look alike, but one may be the only link between others.",
         "If that small link is lost, the whole network can break apart."]
yy = 2.05
for p in probs:
    box(s, 1.55, yy + 0.1, 0.1, 0.1, fill=ORANGE, shape=MSO_SHAPE.OVAL)
    text(s, 1.75, yy, 3.4, 0.7, p, size=12.5)
    yy += 0.75
# objectives card
box(s, 5.55, 1.25, 4.05, 3.15, fill=MINT)
badge(s, 5.75, 1.42, 0.42, "✓", GREEN, size=15)
text(s, 6.3, 1.43, 3.1, 0.45, "Objectives", size=18, bold=True, color=GREEN)
objs = ["Detect mangrove habitat", "Identify habitat patches", "Find important connecting patches",
        "Support protection & restoration decisions"]
yy = 2.0
for o in objs:
    labelled_box(s, 5.8, yy, 3.55, 0.48, o, fill=WHITE, line=GREEN, color=NAVY, size=12.5)
    yy += 0.58
# existing vs our approach
text(s, 1.3, 4.6, 1.9, 0.4, "Existing approach", size=12.5, bold=True, color=MUTED)
labelled_box(s, 3.2, 4.55, 1.7, 0.45, "Satellite map", fill=LIGHT, line=LINE, color=TEXT, size=11.5)
arrow(s, 4.97, 4.66)
labelled_box(s, 5.32, 4.55, 1.9, 0.45, "Where is habitat?", fill=LIGHT, line=LINE, color=TEXT, size=11.5)
text(s, 1.3, 5.25, 1.9, 0.4, "Our approach", size=12.5, bold=True, color=NAVY)
steps = ["Satellite map", "Where is habitat?", "How is it connected?", "Which patches matter?"]
xs = [3.2, 4.73, 6.26, 7.79]
for i, (st, x) in enumerate(zip(steps, xs)):
    last = i == len(steps) - 1
    labelled_box(s, x, 5.2, 1.3 if not last else 1.8, 0.55, st, fill=NAVY if last else WHITE, line=NAVY,
                 color=WHITE if last else NAVY, size=11)
    if not last:
        arrow(s, x + 1.33, 5.37, w=0.2, h=0.2)
s.notes_slide.notes_text_frame.text = (
    "The problem: finding where mangroves are is not enough. Two patches may look the same on a map, but one may "
    "be the only bridge between two groups. Our objective is to detect the habitat, split it into patches, find "
    "the important connecting patches, and help conservation teams decide what to protect or restore.")

# =========================================================================== 4 · DATA + METHODOLOGY
s = new_slide("Data and Methodology")
facts = [("Sentinel-1 satellite radar", "Free ESA images; radar sees through cloud (VV + VH, 10 m)"),
         ("4 coastal study areas", "Vembanad–Kol (Kerala) · Sundarbans (WB) · Gulf of Mannar (TN) · Bhitarkanika (Odisha)"),
         ("Global Mangrove Watch 2020", "Used as reference labels to train the model — not field-verified"),
         ("1,373 image tiles (256 × 256)", "Current run: 800 training · 195 validation · 200 test")]
yy = 1.3
for i, (h, d) in enumerate(facts):
    box(s, 1.3, yy, 4.6, 1.05, fill=LIGHT)
    badge(s, 1.45, yy + 0.3, 0.44, str(i + 1), NAVY, size=13)
    text(s, 2.02, yy + 0.1, 3.8, 0.9, [[(h, {"bold": True, "color": NAVY, "size": 14})], [(d, {"size": 11})]],
         space=1)
    yy += 1.17
flow = [("Satellite images", "Sentinel-1, 2020"), ("Image tiles", "256 × 256 pixels"),
        ("Reference labels", "Global Mangrove Watch"), ("AI model", "learns what mangrove looks like"),
        ("Mangrove probability map", "chance of mangrove per pixel")]
yy = 1.3
for i, (h, d) in enumerate(flow):
    last = i == len(flow) - 1
    labelled_box(s, 6.25, yy, 3.35, 0.72, h, fill=NAVY if last else WHITE, line=NAVY,
                 color=WHITE if last else NAVY, size=13, sub=d, subsize=10)
    if not last:
        arrow(s, 7.8, yy + 0.76, w=0.24, h=0.2, direction="down")
    yy += 0.97
s.notes_slide.notes_text_frame.text = (
    "Our input is Sentinel-1 radar imagery for four coastal areas. For training labels we use Global Mangrove "
    "Watch 2020 — these are reference labels from a published map, not field ground truth. We cut the images into "
    "1,373 tiles; this run used 800 for training, 195 for validation and 200 for testing. The model outputs a "
    "probability of mangrove for every pixel.")

# =========================================================================== 5 · ARCHITECTURE
s = new_slide("System Architecture")
text(s, 1.3, 1.15, 5, 0.35, "How the analysis works", size=14, bold=True, color=NAVY)
pipe = ["Sentinel-1 satellite data", "Preprocessing", "AI mangrove detection",
        "Habitat patch extraction", "Mangrove probability map", "",
        "Connectivity network", "Critical patches + what-if", "Restoration priority"]
# snake layout: row1 L→R, row2 R→L, row3 L→R ; last goes to dashboard
W, H = 2.35, 0.62
cols = [1.3, 4.25, 7.2]
rows = [1.55, 2.6, 3.65]
r1 = ["Sentinel-1 satellite data", "Preprocessing", "AI mangrove detection"]
r2 = ["Mangrove probability map", "Habitat patch extraction", "Connectivity network"]   # drawn right→left
r3 = ["Critical patches + what-if", "Restoration priority", "Web dashboard"]
for i, t in enumerate(r1):
    labelled_box(s, cols[i], rows[0], W, H, t, size=12.5)
    if i < 2: arrow(s, cols[i] + W + 0.18, rows[0] + 0.2, w=0.26)
arrow(s, cols[2] + W / 2 - 0.12, rows[0] + H + 0.1, w=0.24, h=0.26, direction="down")
for i, t in enumerate(r2):             # positions: col2, col1, col0
    labelled_box(s, cols[2 - i], rows[1], W, H, t, size=12.5)
    if i < 2: arrow(s, cols[2 - i] - 0.44, rows[1] + 0.2, w=0.26, direction="left")
arrow(s, cols[0] + W / 2 - 0.12, rows[1] + H + 0.1, w=0.24, h=0.26, direction="down")
for i, t in enumerate(r3):
    last = i == 2
    labelled_box(s, cols[i], rows[2], W, H, t, size=12.5, fill=NAVY if last else WHITE, color=WHITE if last else NAVY)
    if i < 2: arrow(s, cols[i] + W + 0.18, rows[2] + 0.2, w=0.26)
# software stack
text(s, 1.3, 4.55, 5, 0.35, "Software that runs it", size=14, bold=True, color=GREEN)
stack = [("Python + PyTorch", "AI model"), ("FastAPI", "backend server"), ("SQLite", "database"),
         ("Next.js + React", "web frontend"), ("Leaflet + React Flow", "map + network view")]
sw, x = 1.5, 1.3
for i, (h, d) in enumerate(stack):
    labelled_box(s, x, 4.95, sw, 0.78, h, fill=MINT, line=GREEN, color=GREEN, size=11.5, sub=d, subsize=9.5)
    if i < len(stack) - 1: arrow(s, x + sw + 0.03, 5.24, w=0.14, h=0.18, color=GREEN)
    x += sw + 0.2
s.notes_slide.notes_text_frame.text = (
    "Top: the analysis flow. Satellite data is prepared, the AI model detects mangrove and gives a probability "
    "map, we extract habitat patches, connect them into a network, find the critical patches, run what-if tests, "
    "rank restoration options and show everything on a web dashboard. Bottom: the software — Python and PyTorch "
    "for the AI, a FastAPI backend with a SQLite database, and a Next.js/React frontend with a Leaflet map and a "
    "React Flow network view.")

# =========================================================================== 6 · IMPLEMENTATION
s = new_slide("Implementation")
cards = [("Mangrove detection", "A deep-learning model marks mangrove pixels in the image."),
         ("Patch extraction", "Touching mangrove pixels are grouped into habitat patches."),
         ("Connectivity analysis", "Each patch links to its nearest patches (within 5 km)."),
         ("Criticality analysis", "Remove one patch at a time and measure how much connectivity drops."),
         ("Decision support", "Dashboard shows important patches, what-if results and restoration options.")]
yy = 1.25
for i, (h, d) in enumerate(cards, 1):
    box(s, 1.3, yy, 4.55, 0.86, fill=LIGHT if i % 2 else WHITE, line=LINE)
    badge(s, 1.42, yy + 0.2, 0.44, str(i), NAVY, size=13)
    text(s, 2.0, yy + 0.06, 3.8, 0.8, [[(h, {"bold": True, "color": NAVY, "size": 13.5})], [(d, {"size": 11})]],
         space=0)
    yy += 0.96
picture(s, "graph.png", 6.05, 1.25, w=3.55)
text(s, 6.05, 3.46, 3.55, 0.45, "Network view in our app: each circle is a patch, lines are links",
     size=10, color=MUTED, italic=True, align=PP_ALIGN.CENTER)
text(s, 6.05, 4.0, 3.55, 0.35, "Built with", size=12.5, bold=True, color=NAVY)
chips = ["Python", "PyTorch", "FastAPI", "SQLite", "Next.js", "React", "Leaflet", "React Flow"]
cx, cy = 6.05, 4.4
for c in chips:
    w = 0.16 + 0.085 * len(c)
    if cx + w > 9.6:
        cx, cy = 6.05, cy + 0.42
    labelled_box(s, cx, cy, w, 0.34, c, fill=MINT, line=GREEN, color=GREEN, size=10.5)
    cx += w + 0.08
text(s, 6.05, 5.35, 3.55, 0.5, "Model: U-Net with EfficientNet-B0 (≈ 6.25 M parameters)", size=10.5, color=MUTED)
s.notes_slide.notes_text_frame.text = (
    "We implemented five parts. A deep-learning model marks mangrove pixels. Touching pixels are grouped into "
    "patches. Each patch is linked to its nearest patches within 5 km — that forms the network. Then we remove one "
    "patch at a time and measure how much connectivity drops. Finally the dashboard shows the important patches, "
    "the what-if results and restoration options. The model is a U-Net with an EfficientNet-B0 encoder; the larger "
    "B7 version from the research paper is future work.")

# =========================================================================== 7 · RESULTS
s = new_slide("Results and Key Findings")
for x, val, lab, sub in [(1.3, "0.842", "Test IoU", "overlap with reference map"),
                         (3.45, "0.914", "Test F1 score", "balance of precision & recall")]:
    box(s, x, 1.25, 2.0, 1.55, fill=WHITE, line=NAVY, lw=1.75)
    text(s, x, 1.3, 2.0, 0.75, val, size=34, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    text(s, x, 2.02, 2.0, 0.7, [[(lab, {"bold": True, "size": 12.5})], [(sub, {"size": 10, "color": MUTED})]],
         align=PP_ALIGN.CENTER, space=0)
box(s, 5.65, 1.25, 3.95, 1.55, fill=PEACH)
text(s, 5.8, 1.32, 3.7, 1.45, anchor=MSO_ANCHOR.MIDDLE, paras=
     [[("Development model result", {"bold": True, "color": ORANGE, "size": 13})],
      [("Measured on 200 test tiles against Global Mangrove Watch reference labels — not field accuracy. "
        "Strongest on Sundarbans; weaker on thin Kerala mangrove strips.", {"size": 11})]], space=2)
text(s, 1.3, 3.0, 8.3, 0.4, [[("Size alone does not tell us which patch is important", {"bold": True, "color": NAVY,
                                                                                            "size": 16})]])
for x, fill, line, title, rows_, col in [
    (1.3, LIGHT, LINE, "P01 — largest patch", [("35.1 ha", "#1 by size"), ("−30.7%", "connectivity if removed"),
                                              ("No split", "other routes still connect")], NAVY),
    (3.85, PEACH, ORANGE, "P17 — small bridge patch", [("3.1 ha", "only #17 by size"), ("−27.0%", "connectivity if removed"),
                                                       ("2 → 3 groups", "network breaks apart")], ORANGE)]:
    box(s, x, 3.5, 2.4, 2.35, fill=fill, line=line)
    text(s, x + 0.1, 3.57, 2.2, 0.4, title, size=12.5, bold=True, color=col, align=PP_ALIGN.CENTER)
    yy = 4.02
    for big, small in rows_:
        text(s, x + 0.1, yy, 2.2, 0.58, [[(big, {"bold": True, "size": 16, "color": col})],
                                          [(small, {"size": 10, "color": MUTED})]], align=PP_ALIGN.CENTER, space=0)
        yy += 0.6
picture(s, "graph_zoom.png", 7.12, 3.5, h=2.35)
text(s, 6.4, 5.9, 3.2, 0.45, "Network in our app — P17 and P07 are the only “bridge” patches",
     size=10, color=MUTED, italic=True, align=PP_ALIGN.CENTER)
s.notes_slide.notes_text_frame.text = (
    "Our development model reaches a test IoU of 0.842 and an F1 of 0.914 on 200 test tiles. These are measured "
    "against Global Mangrove Watch reference labels, not field surveys, and most of the strength comes from the "
    "Sundarbans. The 95.56% figure in the literature belongs to the foundation paper, not to us. "
    "The key finding: P01 is the biggest patch, but patch P17 is only 3.1 hectares and 17th by size — yet removing "
    "it cuts connectivity by 27% and splits the network into three groups. Size alone does not show importance. "
    "(Kerala 2025 analysis run, development model.)")

# =========================================================================== 8 · WHAT-IF + RESTORATION
s = new_slide("What-if and Restoration")
box(s, 1.3, 1.2, 4.1, 4.65, fill=WHITE, line=LINE)
text(s, 1.45, 1.27, 3.9, 0.45, "What happens if a patch is lost?", size=15, bold=True, color=NAVY)
cd = CategoryChartData()
cd.categories = ["Remove P17 (3.1 ha)", "Remove P01 (35.1 ha)"]
cd.add_series("Habitat lost (%)", (1.4, 16.0))
cd.add_series("Connectivity lost (%)", (27.0, 30.7))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1.4), Inches(1.75), Inches(3.9), Inches(3.05), cd)
ch = gf.chart
ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
ch.legend.font.size = Pt(10); ch.legend.font.name = FONT
for ser, colr in zip(ch.series, [MUTED, ORANGE]):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(colr)
    ser.data_labels.show_value = True; ser.data_labels.number_format = '0.0"%"'
    ser.data_labels.number_format_is_linked = False; ser.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    ser.data_labels.font.size = Pt(10); ser.data_labels.font.bold = True
ch.value_axis.visible = False; ch.value_axis.has_major_gridlines = False
ch.value_axis.maximum_scale = 36
ch.category_axis.tick_labels.font.size = Pt(10); ch.category_axis.tick_labels.font.name = FONT
ch.category_axis.format.line.color.rgb = rgb(LINE)
text(s, 1.45, 4.85, 3.9, 0.95,
     [[("P17 is tiny but costs almost as much connectivity as the biggest patch. ", {"size": 11})],
      [("Simulated scenario under project assumptions.", {"size": 10, "italic": True, "color": MUTED})]], space=1)
box(s, 5.6, 1.2, 4.0, 4.65, fill=WHITE, line=LINE)
text(s, 5.75, 1.27, 3.8, 0.45, "What happens if we restore an area?", size=15, bold=True, color=GREEN)
flow = [("Candidate C1 (1.6 ha)", MINT), ("Added to the network", MINT), ("+1.29% connectivity · 3 new links", GREEN)]
yy = 1.8
for i, (t, f) in enumerate(flow):
    last = i == 2
    labelled_box(s, 5.85, yy, 3.5, 0.42, t, fill=f, line=GREEN, color=WHITE if last else GREEN, size=11.5)
    if not last:
        arrow(s, 7.48, yy + 0.44, w=0.22, h=0.14, direction="down", color=GREEN)
    yy += 0.6
picture(s, "restoration.png", 5.75, 3.62, w=3.7)
text(s, 5.75, 5.12, 3.75, 0.7,
     [[("Potential candidates from model outputs — not field-confirmed. ", {"size": 10})],
      [("No cost data yet, so candidates are ranked by connectivity gain.", {"size": 10, "color": MUTED,
                                                                           "italic": True})]], space=0)
s.notes_slide.notes_text_frame.text = (
    "Left: what if a patch is lost? Removing P17 loses only 1.4% of the habitat but 27% of connectivity. Removing "
    "the largest patch P01 loses 16% of the habitat and 30.7% of connectivity. These are simulations under our "
    "assumptions, not predictions. Right: restoration works the other way — we add a candidate area to the network "
    "and measure the gain. Candidate C1 adds 1.29% connectivity with three new links. Candidates come from model "
    "outputs and are not field-confirmed, and since we have no cost data they are ranked by connectivity gain.")

# =========================================================================== 9 · CONCLUSION
s = new_slide("Conclusion and Future Work")
parts = ["Satellite data", "AI habitat detection", "Connectivity analysis", "What-if scenarios"]
x = 1.3
for i, p in enumerate(parts):
    labelled_box(s, x, 1.3, 1.62, 0.62, p, fill=LIGHT, line=NAVY, size=11.5)
    if i < 3:
        text(s, x + 1.62, 1.33, 0.4, 0.55, "+", size=22, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    x += 2.02
arrow(s, 5.3, 2.02, w=0.26, h=0.26, direction="down")
labelled_box(s, 3.3, 2.35, 4.3, 0.55, "Support for coastal conservation decisions", fill=NAVY, color=WHITE, size=13)
box(s, 1.3, 3.15, 4.05, 2.35, fill=PEACH)
badge(s, 1.45, 3.28, 0.4, "!", ORANGE, size=15)
text(s, 1.95, 3.28, 3.2, 0.42, "Limitations", size=16, bold=True, color=ORANGE)
lims = ["Current model is a development version", "Reference labels are not field-verified ground truth",
        "Field validation and real restoration cost data still needed"]
yy = 3.8
for l in lims:
    box(s, 1.5, yy + 0.1, 0.09, 0.09, fill=ORANGE, shape=MSO_SHAPE.OVAL)
    text(s, 1.68, yy, 3.55, 0.55, l, size=11.5)
    yy += 0.55
box(s, 5.55, 3.15, 4.05, 2.35, fill=MINT)
badge(s, 5.7, 3.28, 0.4, "✓", GREEN, size=14)
text(s, 6.2, 3.28, 3.2, 0.42, "Future Work", size=16, bold=True, color=GREEN)
fut = ["Improve the model for each region (full UNB7 model)", "Add field validation",
       "Use real restoration cost information", "Prepare the system for larger-scale use"]
yy = 3.8
for f in fut:
    box(s, 5.75, yy + 0.1, 0.09, 0.09, fill=GREEN, shape=MSO_SHAPE.OVAL)
    text(s, 5.93, yy, 3.55, 0.45, f, size=11.5)
    yy += 0.42
banner(s, 1.3, 5.68, 8.3, 0.5, [("From mapping habitat to understanding which habitat matters.",
                                  {"bold": True, "italic": True, "color": "FFD54A"})], size=15)
s.notes_slide.notes_text_frame.text = (
    "To conclude: EcoConnectAI combines satellite data, AI habitat detection, connectivity analysis and what-if "
    "scenarios to support coastal conservation decisions. Our limitations: the model is still a development "
    "version, the labels are reference labels rather than field-verified ground truth, and we still need field "
    "validation and real cost data. Next we want to improve the model per region, add field validation, use real "
    "cost information and prepare the system for larger-scale use. From mapping habitat to understanding which "
    "habitat matters. Thank you.")

prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")
